"""
End-to-End Billed Test (Feature 083, Acceptance phase): Fee Agreement Document
Generation.

STATUS (2026-09-13, redesign): rewritten for the "minimal code, maximal AI"
architecture - the AI now composes the ENTIRE document body itself (a tool
call fetches a reference template's body text, the AI writes new full body
text, a thin render tool wraps it in the branded .docx shell) and there is NO
approval gate anywhere in this flow any more (REQ-083-04's self-verification
is the sole release gate, per the original spec - see
specs/repo/features/083-fee-agreement-docs/spec.md). Every prompt in this
file is UNCHANGED from the pre-redesign version (per explicit human
instruction, 2026-09-13: "the tests prompts dont need to change and neither
are the user expectations... The only thing that needs to change is YOUR
IMPLEMENTATION") - only the test MECHANICS changed: no more
pending_local_tool_approval_manager / button-tap-approve round trip; a single
user turn now runs the AI's own get_template -> compose -> render -> verify
-> send loop to completion, and the test observes the outcome by mocking
`WhatsAppHandler.send_document_response` and inspecting the real
`GeneratedDocument` it was called with (which variant was rendered, and the
rendered .docx's actual text).

Tests the real OpenAI function-calling mechanism end-to-end - NOT
unit-testable, since what's under test is whether the real model (a) selects
the right template variant from natural phrasing, (b) composes a correct,
placeholder-free Hebrew document body around the user's actual facts (never
inventing figures), (c) self-verifies before ever sending, and (d) the file
that reaches WhatsApp is intact.

Text-only conversational turns are `billed`; nothing here is `expensive` (no
vision/image calls). The actual WhatsApp `sendFileByUpload` network call is
stubbed at the send boundary only (Gate Zero - a real live send is a
separate, explicitly human-approved step per research.md #2, not exercised
by an automated test run) - this mirrors `test_reminder_lifecycle_billed.py`'s
existing precedent of stubbing `send_proactive_message` for the same reason.
Every other component (AIHandler, DocTemplateEngine, OpenAI) is real, per
CONSTITUTION SS I/SS V.
"""

import logging
import re
import shutil
import uuid
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument
from docx.oxml.ns import qn

from src.models.config import AppConfiguration
from tests.e2e_helpers import sanity_worker_data_root

logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)

GODFATHER_CHAT_ID_TEMPLATE = "{phone}@c.us"


@pytest.mark.billed
class TestFeeAgreementGenerationFlow:
    """Given/When/Then E2E coverage for the UAT stages in
    specs/repo/features/083-fee-agreement-docs/user-stories.md, against the
    2026-09-13 no-approval-gate, AI-authored-body-text architecture.
    """

    @pytest.fixture
    def config(self):
        config_path = Path(__file__).parent.parent.parent / "config" / "config.test.json"
        if not config_path.exists():
            pytest.skip("config.test.json not found")
        config = AppConfiguration.from_file(str(config_path))
        config.validate()
        test_data_root = sanity_worker_data_root()
        config.data_root = str(test_data_root)
        config.memory['session']['storage_dir'] = str(test_data_root / "sessions")
        config.memory['longterm']['storage_dir'] = str(test_data_root / "memory")
        # Feature 083: feature-flagged, default false — must be explicitly on for these tests.
        config.feature_flags = dict(config.feature_flags or {})
        config.feature_flags['fee_agreement_docs'] = True
        return config

    @pytest.fixture
    def denidin_app(self, config):
        import denidin

        if denidin.denidin_app is None:
            config_dict = {
                'green_api_instance_id': config.green_api_instance_id,
                'green_api_token': config.green_api_token,
                'ai_api_key': config.ai_api_key,
                'ai_model': config.ai_model,
                'ai_vision_model': config.ai_vision_model,
                'ai_embedding_model': config.ai_embedding_model,
                'ai_reply_max_tokens': config.ai_reply_max_tokens,
                'log_level': config.log_level,
                'data_root': config.data_root,
                'feature_flags': config.feature_flags,
                'fee_agreements': config.fee_agreements,
                'godfather_phone': config.godfather_phone,
                'memory': config.memory,
                'constitution_config': config.constitution_config,
                'user_roles': config.user_roles,
                'mcp': config.mcp,
            }
            denidin.denidin_app = denidin.initialize_app(config_dict)
        return denidin.denidin_app

    # --- Notification/turn helpers (mirrors test_reminder_lifecycle_billed.py) --------

    @staticmethod
    def _create_notification(chat_id, sender, sender_name, text, msg_id):
        from whatsapp_chatbot_python import Notification

        notification = Notification.__new__(Notification)
        notification.event = {
            'typeWebhook': 'incomingMessageReceived',
            'idMessage': msg_id,
            'senderData': {'chatId': chat_id, 'sender': sender, 'senderName': sender_name},
            'messageData': {'typeMessage': 'textMessage', 'textMessageData': {'textMessage': text}},
        }
        notification._test_sent_messages = []
        notification._test_button_sends = []
        notification._test_document_sends = []

        def track_answer(message):
            notification._test_sent_messages.append(message)
            logger.info(f"Would send to user: {message}")

        def track_answer_with_interactive_buttons(body, buttons, header=None, footer=None):
            from types import SimpleNamespace
            id_message = f"TEST_BUTTONS_{msg_id}_{len(notification._test_button_sends)}"
            notification._test_button_sends.append({'body': body, 'buttons': buttons, 'idMessage': id_message})
            notification._test_sent_messages.append(body)
            return SimpleNamespace(code=200, data={'idMessage': id_message}, error=None)

        notification.answer = track_answer
        notification.answer_with_interactive_buttons = track_answer_with_interactive_buttons
        return notification

    def _send_text(self, chat_id, sender, sender_name, text, label):
        from denidin import handle_text_message
        msg_id = f"billed_{label}_{uuid.uuid4().hex[:8]}"
        notification = self._create_notification(chat_id, sender, sender_name, text, msg_id)
        handle_text_message(notification)
        return notification

    def _godfather(self, config):
        phone = config.godfather_phone
        return phone, GODFATHER_CHAT_ID_TEMPLATE.format(phone=phone)

    @staticmethod
    def _reset_session(denidin_app, chat_id):
        """2026-09-14 (explicit human instruction, following a real observed
        failure): the godfather chat's session is a PERSISTENT on-disk
        record, and `denidin_app` here is a process-wide singleton reused
        across every test in this file (`if denidin.denidin_app is None`) -
        so without an explicit reset, one test's client/date/signer exchange
        leaks into the next test's turn as prior conversation history, and
        the model (correctly) treats it as part of the same ongoing thread
        instead of a fresh request. Deletes the session's own directory and
        its chat_index row directly - SessionManager is a pure read-through
        over chat_index.db + the session directory (no in-memory cache), so
        the next turn for this chat_id transparently creates a brand new
        session."""
        sm = denidin_app.ai_handler.session_manager
        session_id = sm._index_lookup(chat_id) or sm.chat_to_session.get(chat_id)
        if session_id is None:
            return
        session_dir = sm.storage_dir / session_id
        if session_dir.exists():
            shutil.rmtree(session_dir)
        # Must go through the SessionManager's OWN live connection - a
        # separate sqlite3.connect() to the same file would leave
        # `chat_to_session` (the in-memory fast-path cache `get_session()`
        # also falls back to) stale, and get_session() would keep resolving
        # to the just-deleted session_id regardless of what the DB says.
        sm._index_conn.execute("DELETE FROM chat_sessions WHERE chat = ?", (chat_id,))
        sm._index_conn.commit()
        sm.chat_to_session.pop(chat_id, None)

    @pytest.fixture(autouse=True)
    def _isolated_godfather_session(self, denidin_app, config):
        """Wipe the godfather chat's session before AND after every test in
        this file, so each test runs against a clean conversational slate -
        no leakage from a previous test in this run, and no leakage from a
        previous pytest invocation's leftover on-disk session either."""
        phone, chat_id = self._godfather(config)
        self._reset_session(denidin_app, chat_id)
        yield
        self._reset_session(denidin_app, chat_id)

    @staticmethod
    def _last_message(notification):
        return notification._test_sent_messages[-1] if notification._test_sent_messages else None

    @staticmethod
    def _docx_text(path):
        return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)

    # --- Shell/RTL/authored-content assertions (shared by every test below) ------
    #
    # render_free_text() (doc_template_engine.py) only ever touches doc.element.body
    # - it deletes and rebuilds the body's own <w:p> paragraphs from the AI's text,
    # but never touches the header/footer PARTS, which are separate .docx package
    # parts entirely outside the body. So the logo image + footer contact line are
    # STRUCTURALLY guaranteed intact by code, regardless of what the AI wrote - a
    # real regression there (e.g. a future refactor that nukes the wrong XML
    # subtree) would be a code bug, not a model-quality issue. The firm's own name,
    # the document date, and the signature block, by contrast, USED to be constant
    # template body paragraphs (each real template used to have its own "עו"ד אילה
    # הוניגמן"/"תאריך: ___"/"חתימה: ___" lines) but render_free_text() deletes ALL
    # original body paragraphs and replaces them with the AI's own text - so those
    # three are now the AI's own responsibility to include, and are checked as
    # AI-authored-content, not shell integrity.

    @staticmethod
    def _assert_shell_intact(temp_path):
        """(1) Logo/header/footer/RTL - the parts of the .docx code owns, never the
        AI. Every non-empty body paragraph render_free_text() writes must carry the
        same jc="right"/"center"/"both" + paragraph-mark <w:rtl/> + run-level
        <w:rtl/> recipe _build_rtl_paragraph() applies (see doc_template_engine.py)
        - this is what makes a real Word client render right-to-left correctly. A
        paragraph-level <w:bidi/> is scoped exactly to jc="both" paragraphs
        (2026-09-14, two related but opposite real-rendering findings): it actively
        breaks jc="right"/"center" rendering in real Word, so its absence there is
        part of what "intact" means - but jc="both" WITHOUT it was separately shown,
        via a real rendered screenshot, to render a short/single-line justified
        paragraph flush LEFT instead of right, so its PRESENCE there is required."""
        doc = DocxDocument(str(temp_path))
        section = doc.sections[0]

        # Logo: a real image relationship on the header part - untouched by
        # render_free_text(), so its presence proves the branded shell, not the
        # AI's own text, is what's being served.
        header_part = section.header.part
        assert any(rel.reltype.endswith("/image") for rel in header_part.rels.values()), (
            "header logo image is missing - the branded .docx shell was not preserved"
        )

        # Footer: the firm's real contact line is a constant string, identical
        # across all 5 templates - never AI-authored, never a value.
        footer_text = "\n".join(p.text for p in section.footer.paragraphs)
        assert "honigman-law.com" in footer_text, (
            f"footer contact line missing/altered - shell not preserved: {footer_text!r}"
        )

        # RTL: every actual line of AI-authored body text must render right-to-left,
        # with the correct alignment for its jc value - no "left".
        for para in doc.paragraphs:
            if not para.text.strip():
                continue
            pPr = para._p.find(qn('w:pPr'))
            assert pPr is not None, f"paragraph has no pPr (not RTL-safe): {para.text!r}"
            jc = pPr.find(qn('w:jc'))
            jc_val = jc.get(qn('w:val')) if jc is not None else None
            assert jc_val in ('right', 'center', 'both'), (
                f"paragraph alignment is not right/center/both (RTL-safe): "
                f"{jc_val!r} for {para.text!r}"
            )
            mark_rPr = pPr.find(qn('w:rPr'))
            assert mark_rPr is not None and mark_rPr.find(qn('w:rtl')) is not None, (
                f"paragraph mark is not RTL: {para.text!r}"
            )
            for run in para.runs:
                if not run.text.strip():
                    continue
                run_rPr = run._r.find(qn('w:rPr'))
                assert run_rPr is not None and run_rPr.find(qn('w:rtl')) is not None, (
                    f"run text is not RTL: {run.text!r}"
                )
            has_bidi = pPr.find(qn('w:bidi')) is not None
            if jc_val == 'both':
                assert has_bidi, (
                    f"justified (jc=both) paragraph is missing <w:bidi/> - real "
                    f"rendering shows this as flush-LEFT, not right: {para.text!r}"
                )
            else:
                assert not has_bidi, (
                    f"paragraph-level <w:bidi/> present on a non-justified "
                    f"({jc_val!r}) paragraph - this is the confirmed real-Word "
                    f"RTL-rendering bug this feature fixed, must never regress: "
                    f"{para.text!r}"
                )

            # Line spacing: 1.5 (w:line="360" at lineRule="auto", where 240 is
            # single) - 2026-09-14, explicit human instruction: single spacing
            # made a real rendered document look visibly crowded.
            spacing = pPr.find(qn('w:spacing'))
            assert spacing is not None, f"paragraph has no w:spacing: {para.text!r}"
            assert spacing.get(qn('w:line')) == '360', (
                f"paragraph is not 1.5-line-spaced (w:line="
                f"{spacing.get(qn('w:line'))!r}, expected '360'): {para.text!r}"
            )

        # Only the document title may be centered (2026-09-14 fix: "לבין" -
        # the line between the client and the firm in the code-injected
        # header - was previously centered, inconsistent with the
        # party/firm lines around it and the rest of the document).
        centered_texts = [
            p.text.strip() for p in doc.paragraphs
            if p.text.strip()
            and p._p.find(qn('w:pPr')) is not None
            and p._p.find(qn('w:pPr')).find(qn('w:jc')) is not None
            and p._p.find(qn('w:pPr')).find(qn('w:jc')).get(qn('w:val')) == 'center'
        ]
        assert centered_texts == ["הסכם שכר טרחה"], (
            f"expected ONLY the title to be centered, got: {centered_texts!r}"
        )

        # Exactly three sections (S1 header, S2 content, S3 footer - see
        # doc_template_engine.py's DocTemplateEngine class docstring
        # constants), each separated from its neighbor by one full
        # blank-line paragraph (2026-09-14, explicit human instruction) -
        # exactly two such blanks, never zero, never more.
        blank_count = sum(1 for p in doc.paragraphs if not p.text.strip())
        assert blank_count == 2, (
            f"expected exactly 2 section-break blank paragraphs (S1|S2 and "
            f"S2|S3), got {blank_count}: {[p.text for p in doc.paragraphs]!r}"
        )

    @staticmethod
    def _assert_ai_authored_essentials(text, expected_client_name):
        """(2)+(3a) The AI's own authored content must still carry the essentials a
        real fee agreement needs, even though nothing forces this structurally any
        more: the firm/lawyer's own identity, a document date, and a place to sign -
        plus the exact client name the user asked for (not a paraphrase, not a
        different person)."""
        assert expected_client_name in text, (
            f"client name {expected_client_name!r} (exactly as the user gave it) "
            f"not found in the document body: {text!r}"
        )
        assert "הוניגמן" in text, (
            f"the firm/lawyer's own identity (עו\"ד אילה הוניגמן) is missing from "
            f"the document body - this must appear even though it's no longer a "
            f"template constant: {text!r}"
        )
        assert "תאריך" in text or re.search(r"\d{1,2}[./]\d{1,2}[./]\d{2,4}", text), (
            f"no document date (a 'תאריך' label or an actual DD.MM.YYYY-style date) "
            f"found anywhere in the body: {text!r}"
        )
        # 2026-09-14: "חתימ" (not "חתימה" alone) so a legitimate real phrasing
        # variant like "חתימת הלקוח" (found in a real billed run) still
        # matches - "חתימה"/"חתימת"/"לחתום" all share this root.
        assert any(word in text for word in ("חתימ", "החתום", "ולראיה")), (
            f"no signature block/place-to-sign found anywhere in the body: {text!r}"
        )

    # --- Stage 1: Template Selection Accuracy + self-verification gate ------------

    @pytest.mark.parametrize(
        "user_text,followup_texts,expected_variant,expected_client_name",
        [
            (
                # "A simple, regular agreement" (explicit human framing, 2026-09-12):
                # NOT hourly_consultation, NOT alternative_tracks - a plain
                # single-fee-for-a-defined-scope request. Per the 2026-09-14
                # variant-count reduction (the real corpus had zero examples
                # of a genuinely standalone "fixed_price_project" template -
                # see config/fee_agreement_templates/examples/README.md), this
                # shape is now multi_component_agreement with exactly ONE
                # component, not a separate variant.
                # 2026-09-14: everything the AI could plausibly need to ask
                # about (date, signer) is given up front in this single
                # message, per explicit human instruction - this case proves
                # the one-shot no-approval flow when the user front-loads
                # every detail themselves.
                "תכין הסכם שכר טרחה רגיל עבור מר אריאל בכר, בנושא בניית אתר "
                "אינטרנט. שכר הטרחה 25,000 ש\"ח כולל מע\"מ, לתשלום תוך 60 "
                "יום. התאריך: היום. החתימה מטעם הלקוח תהיה של מר אריאל בכר "
                "עצמו.",
                [],
                "multi_component_agreement",
                "אריאל בכר",
            ),
            (
                # 2026-09-14: deliberately terse up front (per explicit human
                # instruction) - this case proves the AI's own clarifying
                # questions (date, signer, VAT treatment) get answered over
                # the course of the conversation instead, still ending in the
                # same one-shot send once every detail is in. "תכין הסכם" is
                # included up front (unlike the plain facts-only phrasing
                # first tried here) - without it the model correctly reads
                # this as dictating ledger facts to record, not a document
                # request, and never touches the fee-agreement tools at all.
                # "יוסי זאנזן" is a real, pre-seeded Morning sandbox client
                # (tests/fixtures/morning_sandbox_clients.json, seeded once
                # via a real add_client conversational turn, 2026-09-14) -
                # exact-name resolution, no ambiguous-candidate detour.
                "תכין הסכם שכר טרחה עבור יוסי זאנזן, 10000 צו מניעה, 25% מזכיה",
                [
                    "הסכום כולל מע\"מ. התאריך: היום. החתימה מטעם הלקוח תהיה "
                    "של יוסי זאנזן עצמו.",
                ],
                "multi_component_agreement",
                "יוסי זאנזן",
            ),
        ],
        ids=["simple_regular_agreement", "multi_component_agreement"],
    )
    def test_stage1_template_selection(
        self, denidin_app, config, user_text, followup_texts, expected_variant,
        expected_client_name
    ):
        """The AI selects the correct template variant from phrasing alone,
        composes a full document body around it, self-verifies (no leftover
        {{PLACEHOLDER}} tokens), and only then reaches the send boundary - all
        within one turn (or a short natural back-and-forth answering the AI's
        own clarifying questions), no human approval step. Surfaced by
        inspecting the real GeneratedDocument the (mocked) send call
        received. Also asserts the branded shell (logo/footer/RTL) was
        preserved intact and the AI's own authored content (client name,
        firm identity, date, signature block) is present."""
        phone, chat_id = self._godfather(config)

        # send_document_response(self, generated, chat_id, caption) is patched at the
        # class level, so the mock receives no `self` - `generated` is args[0]. Capture
        # the temp file's own path/existence/text INSIDE the mock (2026-09-14 fix) - the
        # real handler unlinks the temp file in a `finally` immediately after this call
        # returns (SC-003 cleanup, always, regardless of send outcome), so reading it
        # back afterward races a file that's already gone by the time the turn completes.
        captured = {}

        def _capture(generated, chat_id=None, caption=None):  # pylint: disable=unused-argument
            captured["variant_id"] = generated.variant_id
            captured["verified"] = generated.verified
            temp_path = Path(generated.temp_path)
            captured["existed_at_send"] = temp_path.exists()
            if temp_path.exists():
                captured["text"] = self._docx_text(temp_path)
                try:
                    self._assert_shell_intact(temp_path)
                    captured["shell_ok"] = True
                except AssertionError as e:
                    captured["shell_ok"] = False
                    captured["shell_error"] = str(e)
            return True

        with patch(
            "src.handlers.whatsapp_handler.WhatsAppHandler.send_document_response",
            side_effect=_capture,
        ) as mock_send:
            self._send_text(chat_id, phone, "Test Godfather", user_text, "stage1")
            for i, followup in enumerate(followup_texts):
                self._send_text(chat_id, phone, "Test Godfather", followup, f"stage1_followup{i}")

            assert mock_send.called, (
                f"expected the document-send boundary to be reached for {user_text!r} "
                f"(self-verification must have passed for this to happen at all)"
            )
            assert captured["variant_id"] == expected_variant, (
                f"expected variant {expected_variant!r}, got {captured['variant_id']!r}"
            )
            assert captured["verified"] is True, (
                "a document reaching send_document_response() must have verified=True — "
                "code-level guard, not just a prompt-level expectation"
            )
            assert captured["existed_at_send"], (
                "temp file must still exist at the moment of send (before cleanup)"
            )
            assert captured["shell_ok"], (
                f"branded shell not intact at send time: {captured.get('shell_error')}"
            )
            text = captured["text"]
            assert "{{" not in text, f"no unresolved {{{{PLACEHOLDER}}}} tokens may remain: {text!r}"
            self._assert_ai_authored_essentials(text, expected_client_name)

    # --- Stage 1b: multi-component supports ANY N>1 (data-model.md "Variable-length
    # component rows") - not a fixed cap ------------------------------------------

    @pytest.mark.parametrize("n_components", [2, 4])
    def test_multi_component_arbitrary_n(self, denidin_app, config, n_components):
        """Per explicit human correction (2026-09-12): the multi-component variant
        must handle ANY N>1 real components, not a fixed maximum - all composed by
        the AI directly into the document body text now (no repeating-table-row
        mechanism), so verified by counting how many of the given fee components
        actually appear, verbatim, in the rendered document."""
        phone, chat_id = self._godfather(config)
        component_descs = [
            "הגשת התביעה - 5,000 ש\"ח כולל מע\"מ",
            "דיון הוכחות אם יידרש - 4,000 ש\"ח כולל מע\"מ",
            "ניהול ההליך עד להחלטה - 8,000 ש\"ח כולל מע\"מ",
            "ערעור אם יוגש - 6,000 ש\"ח כולל מע\"מ",
        ][:n_components]

        with patch(
            "src.handlers.whatsapp_handler.WhatsAppHandler.send_document_response",
            return_value=True,
        ) as mock_send:
            self._send_text(
                chat_id, phone, "Test Godfather",
                "תכין הסכם עבור דוד כרמלי, בתביעה כספית נגד שכנו. שכר הטרחה: "
                + "; ".join(component_descs) + ".",
                f"stage1b-{n_components}",
            )

            assert mock_send.called
            sent_document = mock_send.call_args.args[0] if mock_send.call_args.args \
                else mock_send.call_args.kwargs.get("generated")
            assert sent_document is not None and sent_document.verified is True
            self._assert_shell_intact(sent_document.temp_path)
            text = self._docx_text(sent_document.temp_path)
            assert "{{" not in text, f"leftover placeholder in body: {text!r}"
            self._assert_ai_authored_essentials(text, "דוד כרמלי")
            # Every component's fee amount must appear verbatim - nothing merged
            # or dropped as the count grows.
            for amount in ("5,000", "4,000", "8,000", "6,000")[:n_components]:
                assert amount in text, f"expected amount {amount!r} in body: {text!r}"

    # --- Stage 1c: alternative_tracks - a real choice between mutually-exclusive
    # fee structures, distinct from multi_component_agreement's "all apply together"
    # shape (corpus-driven 5th variant) --------------------------------------------

    def test_alternative_tracks_selected_and_generated(self, denidin_app, config):
        """A request describing two or more mutually-exclusive fee tracks for the
        SAME engagement (the client picks exactly ONE) must select
        `alternative_tracks`, never `multi_component_agreement` (where every
        component applies together)."""
        phone, chat_id = self._godfather(config)

        with patch(
            "src.handlers.whatsapp_handler.WhatsAppHandler.send_document_response",
            return_value=True,
        ) as mock_send:
            self._send_text(
                chat_id, phone, "Test Godfather",
                "תכין הסכם עבור רונית אשכנזי, "
                "עם שני מסלולי שכר טרחה חלופיים לבחירתה - רק מסלול אחד בפועל "
                "יחול: מסלול א' - שכר טרחה קבוע של 18,000 ש\"ח כולל מע\"מ, ללא תלות "
                "בתוצאה. מסלול ב' - שכר טרחה מוזל בסך 10,000 ש\"ח כולל מע\"מ בתוספת "
                "7% מהסכום שייפסק או ייגבה. אין תוספות החלות על שני המסלולים.",
                "alt_tracks_1",
            )

            assert mock_send.called
            sent_document = mock_send.call_args.args[0] if mock_send.call_args.args \
                else mock_send.call_args.kwargs.get("generated")
            assert sent_document is not None
            assert sent_document.variant_id == "alternative_tracks", (
                f"a mutually-exclusive CHOICE between fee structures must select "
                f"alternative_tracks, not {sent_document.variant_id!r} - "
                f"multi_component_agreement is for components that ALL apply together"
            )
            assert sent_document.verified is True

            self._assert_shell_intact(sent_document.temp_path)
            text = self._docx_text(sent_document.temp_path)
            assert "{{" not in text, f"leftover placeholder token(s) found: {text!r}"
            self._assert_ai_authored_essentials(text, "רונית אשכנזי")
            assert "18,000" in text or "18000" in text
            assert "10,000" in text or "10000" in text
            assert "7" in text and "%" in text

    # --- Stage 4: Successful Delivery ---------------------------------------------

    def test_stage4_delivery_and_cleanup(self, denidin_app, config):
        """The Green API sendFileByUpload-backed send path must be invoked exactly
        once with the real generated file, and the temp file must be deleted
        afterward (SC-003 — no leaked files) - all within one turn, no approval
        step. Content/shell is captured inside the mock's side_effect, at the
        moment of the call - the temp file is deleted immediately afterward as
        part of cleanup, so it can't be read back from the assertions below."""
        phone, chat_id = self._godfather(config)

        captured = {}

        def _capture_and_stub(*args, **kwargs):
            path = Path(kwargs.get("path") or args[1])
            captured["path"] = path
            captured["shell_ok"] = True
            try:
                self._assert_shell_intact(path)
            except AssertionError as e:
                captured["shell_ok"] = False
                captured["shell_error"] = str(e)
            captured["text"] = self._docx_text(path)
            return None

        with patch(
            "src.handlers.whatsapp_handler.WhatsAppHandler._send_file_with_retry",
            side_effect=_capture_and_stub,
        ) as mock_upload:
            self._send_text(
                chat_id, phone, "Test Godfather",
                "תכין הסכם שכר טרחה שעתי עם חברת בטא בע\"מ, לצורך ליווי משפטי "
                "שוטף למר יעקב שני, מנכ\"ל החברה, לפי שעות עבודה. תעריף השעה "
                "500 ש\"ח כולל מע\"מ, עד לתקרה של 10,000 ש\"ח.",
                "stage4a",
            )

            assert mock_upload.call_count == 1, (
                f"expected exactly one sendFileByUpload-equivalent call, "
                f"got {mock_upload.call_count}"
            )
            sent_path = captured["path"]
            assert sent_path.suffix == ".docx"
            assert captured["shell_ok"], (
                f"branded shell not intact at send time: {captured.get('shell_error')}"
            )
            text = captured["text"]
            assert "{{" not in text, f"no unresolved placeholder tokens may remain: {text!r}"
            # The contact person named in the prompt is Mr. Yaakov Shani, not the
            # company itself - the AI's own choice of which name(s) to write is not
            # constrained here, only that the actual essentials are present.
            self._assert_ai_authored_essentials(text, "בטא")

            # The boundary call happens BEFORE cleanup — the temp file must be gone
            # by the time the turn has fully completed (checked here, after).
            assert not sent_path.exists(), (
                "temp .docx must be deleted after a successful send (SC-003) — "
                "found still on disk after the turn completed"
            )

    # --- Component terms: example-pattern-guided AND free-form creative ----------

    def test_multi_component_percentage_coshare_payer_terms(self, denidin_app, config):
        """Per human feedback (2026-09-12): components must support percentages,
        cost-sharing with other partners, and a payer entity other than the Client -
        composed as free text by the AI directly into the document body. One
        component here matches a familiar pattern (percentage split with a named
        partner); another describes a genuinely novel arrangement, to prove the AI
        isn't forced to distort it into a nearest-match pattern."""
        phone, chat_id = self._godfather(config)

        with patch(
            "src.handlers.whatsapp_handler.WhatsAppHandler.send_document_response",
            return_value=True,
        ) as mock_send:
            self._send_text(
                chat_id, phone, "Test Godfather",
                "תכין הסכם עבור אבינועם שגיא, "
                "עם שני רכיבי שכר טרחה: (1) דמי תיווך בשיעור 15% מהסכום שייגבה, "
                "בחלוקה 50/50 עם השותפה עו\"ד רותם לוי, לתשלום על ידי הלקוח עם קבלת "
                "הכסף; (2) בונוס הצלחה חד-פעמי בסך 10,000 ש\"ח כולל מע\"מ, לתשלום "
                "ישירות על ידי חברת גורן נכסים בע\"מ שבבעלותו, רק אם העסקה תיסגר "
                "לפני סוף השנה. שכר הטרחה הכולל: כאמור לעיל.",
                "creative_terms",
            )

            assert mock_send.called
            sent_document = mock_send.call_args.args[0] if mock_send.call_args.args \
                else mock_send.call_args.kwargs.get("generated")
            assert sent_document is not None and sent_document.verified is True
            self._assert_shell_intact(sent_document.temp_path)
            text = self._docx_text(sent_document.temp_path)
            self._assert_ai_authored_essentials(text, "אבינועם שגיא")

            # Component 1: percentage + cost-share (matches a familiar pattern).
            assert "15" in text and "%" in text
            assert "רותם לוי" in text
            assert "50" in text  # the split ratio, stated verbatim, not invented
            # Component 2: a non-Client payer entity AND a conditional trigger.
            assert "10,000" in text or "10000" in text
            assert "גורן נכסים" in text
            assert "סוף השנה" in text or "השנה" in text
            # Nothing invented: no percentage/split/payer/condition appears that
            # wasn't actually stated above (spot-check a plausible hallucination).
            assert "20%" not in text and "30,000" not in text and "3000" not in text
