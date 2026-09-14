"""
DocTemplateEngine - Fee Agreement Document Generation (Feature 083).

Fills a .docx template variant's placeholders with AI-supplied values and
writes a temporary output file, without ever inventing a value the AI did not
explicitly provide (REQ-083-02). See
specs/repo/features/083-fee-agreement-docs/contracts/doc-template-engine.md
for the full contract this class implements.

Uses python-docx for both loading/replacing (generation direction) and
read-back (verification direction) - the same library DOCXExtractor already
depends on for the reverse (extraction) direction, per research.md #6.
"""

import copy
import json
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from src.models.fee_agreement import FeeAgreementVariant, GeneratedDocument
from src.utils.logger import get_logger
from src.utils.time_utils import now_local

logger = get_logger(__name__)

PLACEHOLDER_PATTERN = re.compile(r"\{\{[A-Z_]+\}\}")


class DocTemplateEngine:
    """Loads fee agreement template variants and fills them with AI-supplied
    values. Never guesses/defaults a missing value - every validation failure
    is raised as a ValueError so the caller (ai_handler.py's tool dispatch)
    can surface it back to the model as a tool-call error, per contract."""

    def __init__(self, templates_dir: Path, tmp_dir: Path) -> None:
        self.templates_dir = Path(templates_dir)
        self.tmp_dir = Path(tmp_dir)
        self._variants: Optional[Dict[str, FeeAgreementVariant]] = None

    def list_variants(self) -> List[FeeAgreementVariant]:
        """Loads/returns manifest.json's variants. Read-only; no side effects."""
        return list(self._load_variants().values())

    def _load_variants(self) -> Dict[str, FeeAgreementVariant]:
        if self._variants is None:
            manifest_path = self.templates_dir / "manifest.json"
            with open(manifest_path, "r", encoding="utf-8") as f:
                manifest = json.load(f)
            self._variants = {
                v["variant_id"]: FeeAgreementVariant.from_dict(v)
                for v in manifest["variants"]
            }
        return self._variants

    def _get_variant(self, variant_id: str) -> FeeAgreementVariant:
        variants = self._load_variants()
        if variant_id not in variants:
            raise ValueError(
                f"Unknown fee agreement variant_id: {variant_id!r}. "
                f"Known variants: {sorted(variants.keys())}"
            )
        return variants[variant_id]

    def generate(
        self,
        variant_id: str,
        values: Dict[str, str],
        components: Optional[List[Dict[str, str]]] = None,
    ) -> GeneratedDocument:
        variant = self._get_variant(variant_id)
        self._validate_values(variant, values)
        self._validate_components(variant, components)

        doc = Document(str(self.templates_dir / variant.template_filename))
        self._replace_scalar_placeholders(doc, values)
        if variant.repeating_group is not None:
            self._clone_repeating_group(doc, variant.repeating_group, components or [])

        self.tmp_dir.mkdir(parents=True, exist_ok=True)
        document_id = str(uuid.uuid4())
        temp_path = self.tmp_dir / f"{document_id}.docx"
        doc.save(str(temp_path))

        return GeneratedDocument(
            document_id=document_id,
            variant_id=variant_id,
            values=dict(values),
            temp_path=temp_path,
            created_at=now_local(),
            components=list(components) if components else None,
            verified=False,
            sent=False,
        )

    def verify(self, generated: GeneratedDocument) -> Dict[str, Any]:
        """Reads the generated document back and reports whether it's clean.
        Returns the raw facts (extracted text, leftover placeholder tokens,
        which expected values were/weren't found verbatim) - the ACTUAL
        accept/reject judgment is the model's, per REQ-083-04. This method
        never itself decides pass/fail.

        2026-09-13 redesign: a document produced by render_free_text has no
        fixed `values`/`components` schema to cross-check against (the AI
        wrote the whole body itself) - for those, the only fact worth
        reporting is whether any literal "{{...}}" placeholder-looking token
        leaked into the output (the AI should never emit one; the shell's
        own header/footer never contain one either)."""
        doc = Document(str(generated.temp_path))
        full_text = self._extract_full_text(doc)

        remaining_placeholders = sorted(set(PLACEHOLDER_PATTERN.findall(full_text)))
        # 2026-09-14 (explicit human instruction): a real fact, never a
        # code-level gate - REQ-083-04 still leaves accept/reject to the
        # model. python-docx has no layout engine and can never itself know
        # how a .docx paginates in real Word, so this is a real (LibreOffice
        # headless -> PDF -> PyMuPDF) page count, not an estimate. None means
        # the check itself couldn't run (e.g. LibreOffice unavailable in this
        # environment) - the model should not treat None as "1 page, fine".
        page_count = self.count_pages(generated.temp_path)

        if generated.body_text is not None:
            return {
                "document_id": generated.document_id,
                "extracted_text": full_text,
                "remaining_placeholders": remaining_placeholders,
                "page_count": page_count,
                "clean": not remaining_placeholders,
            }

        expected_values = list(generated.values.values())
        if generated.components:
            for entry in generated.components:
                expected_values.extend(entry.values())

        missing_values = [v for v in expected_values if v.strip() and v not in full_text]

        return {
            "document_id": generated.document_id,
            "page_count": page_count,
            "extracted_text": full_text,
            "remaining_placeholders": remaining_placeholders,
            "missing_values": missing_values,
            "clean": not remaining_placeholders and not missing_values,
        }

    # -- free-text authorship (2026-09-13 redesign) ----------------------

    def get_reference_body(self, variant_id: str) -> str:
        """Returns the variant's current body text (placeholders and all) as
        a REFERENCE for the AI to pattern its own, freely-composed body on -
        not something to fill in or reuse verbatim. The firm's letterhead
        (logo/header/footer) is not included here - it lives in the shell
        and is applied automatically by render_free_text, never something
        the AI writes or sees as editable text."""
        variant = self._get_variant(variant_id)
        doc = Document(str(self.templates_dir / variant.template_filename))
        return "\n".join(p.text for p in doc.paragraphs)

    def get_reference_materials(self, variant_id: str) -> Dict[str, Any]:
        """2026-09-14 follow-up: get_reference_body() alone gives the AI only
        ONE example (the template's own placeholder-laden skeleton) to infer
        professional-grade Hebrew legal phrasing from - too thin a basis to
        reliably generalize a firm's actual register/tone. This method adds,
        alongside that one skeleton, a curated set of REAL fee-agreement
        excerpts this firm has actually sent (names/adversaries/exact amounts
        obfuscated - see config/fee_agreement_templates/examples/README.md)
        plus a directive prompt explaining what the AI is meant to do with
        both. Returns {"template_body", "examples", "directive"} -
        `examples`/`directive` default to an empty list / a generic fallback
        instruction when no curated examples file exists yet for this variant
        (never an error - the reference template body is always the floor)."""
        template_body = self.get_reference_body(variant_id)
        examples_path = self.templates_dir / "examples" / f"{variant_id}.json"
        if examples_path.exists():
            with open(examples_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            examples = list(data.get("examples", []))
            directive = data.get("directive") or self._default_directive()
        else:
            examples = []
            directive = self._default_directive()
        return {
            "template_body": template_body,
            "examples": examples,
            "directive": directive,
        }

    @staticmethod
    def _default_directive() -> str:
        return (
            "No curated real-world examples exist yet for this variant - use "
            "the reference template body above as your only style guide. "
            "Write in the same professional Hebrew legal register: formal, "
            "precise, no filler, matching its structure and tone, filled with "
            "the real facts from this conversation only."
        )

    # Constant firm identity (2026-09-14, explicit human instruction: "firm
    # identity is CONSTANT! IT NEVER CHANGES!!" - a real billed run had the AI
    # write "המשרד" instead of naming the firm, which self-verification has
    # no way to catch since it's a fact only a human/code, not the model's
    # memory, can guarantee). Injected by CODE into every rendered document -
    # never something the AI is trusted to remember to write.
    FIRM_LAWYER_NAME = 'עו"ד אילה הוניגמן'
    _TITLE_TEXT = "הסכם שכר טרחה"
    _SIGNATURE_LINE = "____________________"

    def render_free_text(
        self, variant_id: str, client_name: str, body_text: str
    ) -> GeneratedDocument:
        """Replaces the template's ENTIRE body, but NOT with the AI's text
        alone: the constant boilerplate every fee agreement must always
        carry (title, date, firm identity, signature block) is CODE-OWNED
        and injected here, verbatim, every time - never left to the AI's own
        recall. `body_text` is only the substantive, per-document content
        (scope of work, fee terms, conditions) that sits between that fixed
        header and fixed footer. `client_name` is a separate, explicit field
        (not buried inside body_text) so the header block can name the real
        client with code, not a paraphrase the AI might drift on.

        Minimal-code redesign otherwise unchanged (2026-09-13): the AI still
        has full authorship over the substantive clauses' wording, structure,
        and numbering - only the boilerplate around it is fixed. Light
        markup in `body_text` (2026-09-14, visual-fidelity fix): a line
        starting with "## " renders as a bold section heading; "**...**"
        spans within any line render bold - the ONLY formatting vocabulary
        the AI has, matching how a real Word agreement actually looks
        (headings, emphasis) instead of one flat run of plain text per line."""
        if not isinstance(body_text, str) or not body_text.strip():
            raise ValueError("body_text must be a non-empty string")
        if not isinstance(client_name, str) or not client_name.strip():
            raise ValueError("client_name must be a non-empty string")

        variant = self._get_variant(variant_id)
        doc = Document(str(self.templates_dir / variant.template_filename))
        body = doc.element.body

        for p in list(doc.paragraphs):
            p._p.getparent().remove(p._p)
        # 2026-09-14 bug fix: some older templates (pre-dating the "AI writes
        # the whole body" redesign) still carry a native Word TABLE for their
        # old fixed-schema repeating group (e.g. multi_component_agreement's
        # רכיב/תנאים component rows), with its own placeholder cells
        # ({{COMPONENT_LABEL}}/{{COMPONENT_TERMS}}). A <w:tbl> is a sibling
        # of <w:p> in the body, never a paragraph itself, so the loop above
        # never touched it - it survived into every rendered document
        # regardless of the AI's own body_text, permanently tripping
        # verify()'s placeholder check (found via a real billed-test run,
        # 2026-09-14: the model correctly saw "clean: false", couldn't fix a
        # table it has no handle on, and gave up on this variant entirely).
        # The AI owns the substantive body now, tables included if it wants
        # one (as literal text) - no template table may survive rendering.
        for table in list(doc.tables):
            table._tbl.getparent().remove(table._tbl)

        sect_pr = body.find(qn("w:sectPr"))

        def _insert(new_p):
            if sect_pr is not None:
                sect_pr.addprevious(new_p)
            else:
                body.append(new_p)

        today = now_local().strftime("%d.%m.%Y")

        # -- code-owned header block (title, date, identity) - compact by
        # design (few lines, tight spacing) since a real agreement like this
        # must fit on one page (2026-09-14 instruction). --
        _insert(self._build_rtl_paragraph(
            self._TITLE_TEXT, bold=True, size=28, center=True, space_after=80
        ))
        _insert(self._build_rtl_paragraph(f"תאריך: {today}", space_after=80))
        _insert(self._build_rtl_paragraph(
            f'בין {client_name} (להלן – הלקוח) לבין {self.FIRM_LAWYER_NAME} '
            f'(להלן – עוה"ד)',
            space_after=160,
        ))

        # -- the AI's own substantive content --
        # Blank lines (paragraph breaks in the AI's own prose) are never
        # rendered as their own empty paragraph - each one still carries a
        # non-empty paragraph's own space_after PLUS its own, stacking into
        # a visibly oversized gap between sections (found 2026-09-14 via a
        # real screenshot). Spacing between sections is already handled by
        # each paragraph's own space_after; a blank line in the input adds
        # nothing further.
        for line in body_text.split("\n"):
            if not line.strip():
                continue
            if line.startswith("## "):
                _insert(self._build_rtl_paragraph(
                    line[3:], bold=True, size=24, space_after=100
                ))
            else:
                _insert(self._build_rtl_paragraph(line, space_after=80))

        # -- code-owned footer block (signature) - always present, always
        # names the real client, never left to the AI to remember. --
        _insert(self._build_rtl_paragraph(
            "אני מאשר את ההסכם.", space_after=80
        ))
        _insert(self._build_rtl_paragraph(
            f"תאריך: {today}      שם הלקוח: {client_name}      "
            f"חתימה: {self._SIGNATURE_LINE}",
            space_after=0,
        ))

        self.tmp_dir.mkdir(parents=True, exist_ok=True)
        document_id = str(uuid.uuid4())
        temp_path = self.tmp_dir / f"{document_id}.docx"
        doc.save(str(temp_path))

        return GeneratedDocument(
            document_id=document_id,
            variant_id=variant_id,
            values={"client_name": client_name},
            temp_path=temp_path,
            created_at=now_local(),
            body_text=body_text,
            verified=False,
            sent=False,
        )

    @classmethod
    def _build_rtl_paragraph(
        cls,
        text: str,
        *,
        bold: bool = False,
        size: Optional[int] = None,
        center: bool = False,
        space_after: int = 120,
    ):
        """Builds a right-aligned (or centered, for the title) Hebrew
        paragraph, matching the EXACT structure confirmed (2026-09-13, by
        diffing a real human-verified-working .docx) to actually render
        right-to-left in real Word: <w:rtl/> on the run AND on the paragraph
        mark's own rPr - deliberately NO paragraph-level <w:bidi/>, which was
        proven (by that same diff) to break jc="right" rendering.

        2026-09-14 visual-fidelity fix: real agreements have headings,
        emphasis, and deliberate compact spacing - a flat, uniform run of
        plain 11pt text per line (the original redesign) looked nothing like
        the genuine template it replaced and needlessly ran to 2 pages. This
        now supports a per-paragraph size/bold/center/spacing, plus inline
        "**bold**" spans within `text` (the only markup vocabulary the AI is
        given - see render_free_text's docstring). `space_after` is in
        twentieths of a point (Word's own unit) - the default (120 = 6pt) is
        deliberately tighter than Word's own default (~10pt) to help a
        real agreement's worth of text actually fit on one page."""
        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center" if center else "right")
        pPr.append(jc)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), str(space_after))
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")
        pPr.append(spacing)
        # OOXML's CT_RPr schema is a strict, ordered sequence (b/bCs before
        # sz/szCs before rtl, per ECMA-376) - Word/LibreOffice silently drop
        # properties that appear out of order rather than erroring, which is
        # part of what made every "bold" heading/span render as plain text
        # despite <w:b/> being present in the XML (found 2026-09-14 via a
        # real screenshot of Word's actual rendering). b/bCs/sz must come
        # BEFORE rtl. The other, bigger part: for complex-script text (Hebrew/
        # RTL, i.e. any run carrying <w:rtl/>), Word/LibreOffice render bold
        # based on <w:bCs/> (bold complex-script), NOT <w:b/> alone - <w:b/>
        # governs only the Latin/ASCII font. Every bold run here needs BOTH.
        mark_rPr = OxmlElement("w:rPr")
        if bold:
            mark_rPr.append(OxmlElement("w:b"))
            mark_rPr.append(OxmlElement("w:bCs"))
        if size is not None:
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), str(size))
            mark_rPr.append(sz)
            szCs = OxmlElement("w:szCs")
            szCs.set(qn("w:val"), str(size))
            mark_rPr.append(szCs)
        mark_rPr.append(OxmlElement("w:rtl"))
        pPr.append(mark_rPr)
        p.append(pPr)

        for span_text, span_bold in cls._split_bold_spans(text):
            if not span_text:
                continue
            r = OxmlElement("w:r")
            run_rPr = OxmlElement("w:rPr")
            if bold or span_bold:
                run_rPr.append(OxmlElement("w:b"))
                run_rPr.append(OxmlElement("w:bCs"))
            if size is not None:
                sz = OxmlElement("w:sz")
                sz.set(qn("w:val"), str(size))
                run_rPr.append(sz)
                szCs = OxmlElement("w:szCs")
                szCs.set(qn("w:val"), str(size))
                run_rPr.append(szCs)
            run_rPr.append(OxmlElement("w:rtl"))
            r.append(run_rPr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = span_text
            r.append(t)
            p.append(r)
        return p

    @staticmethod
    def _split_bold_spans(text: str):
        """Splits `text` on "**...**" markers into (span_text, is_bold)
        pairs - the AI's only inline-emphasis vocabulary (see
        render_free_text's docstring). An unpaired "**" is treated as plain
        literal text, never an error - the AI's prose should never crash
        rendering over a stray marker."""
        parts = re.split(r"\*\*(.+?)\*\*", text)
        # re.split with one capturing group alternates: [plain, bold, plain,
        # bold, ..., plain] - even indices are plain text, odd are bold.
        return [(part, idx % 2 == 1) for idx, part in enumerate(parts)]

    # -- validation -----------------------------------------------------

    @staticmethod
    def _validate_values(variant: FeeAgreementVariant, values: Dict[str, str]) -> None:
        expected = set(variant.placeholders)
        got = set(values.keys())
        if got != expected:
            missing = expected - got
            extra = got - expected
            raise ValueError(
                f"values keys mismatch for variant {variant.variant_id!r}: "
                f"missing={sorted(missing)} extra={sorted(extra)}"
            )
        for key, value in values.items():
            if not isinstance(value, str) or not value.strip():
                raise ValueError(
                    f"values[{key!r}] must be a non-empty string for variant "
                    f"{variant.variant_id!r} (REQ-083-02: empty is not a legitimate answer)"
                )

    @staticmethod
    def _validate_components(
        variant: FeeAgreementVariant, components: Optional[List[Dict[str, str]]]
    ) -> None:
        rg = variant.repeating_group
        if rg is None:
            if components is not None:
                raise ValueError(
                    f"variant {variant.variant_id!r} has no repeating_group - "
                    f"`components` must not be supplied"
                )
            return

        if not components or len(components) < rg.min_items:
            raise ValueError(
                f"variant {variant.variant_id!r} requires a `components` list of at "
                f"least {rg.min_items} entries; got "
                f"{0 if not components else len(components)}"
            )
        expected_keys = set(rg.row_placeholders)
        # components entries use lowercase tool-schema keys (label/terms) mapped
        # 1:1 in order onto row_placeholders (e.g. COMPONENT_LABEL/COMPONENT_TERMS,
        # TRACK_LABEL/TRACK_TERMS) - the tool schema always uses {label, terms}.
        for entry in components:
            if set(entry.keys()) != {"label", "terms"}:
                raise ValueError(
                    f"each components entry for variant {variant.variant_id!r} must have "
                    f"exactly keys 'label' and 'terms'; got {sorted(entry.keys())}"
                )
            for key in ("label", "terms"):
                if not isinstance(entry[key], str) or not entry[key].strip():
                    raise ValueError(
                        f"components entry {key!r} must be a non-empty string "
                        f"(REQ-083-02: empty is not a legitimate answer)"
                    )
        del expected_keys  # documents the row_placeholders <-> label/terms mapping above

    # -- replacement ------------------------------------------------------

    @staticmethod
    def _iter_all_paragraphs(doc: Document):
        for p in doc.paragraphs:
            yield p
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        yield p

    def _replace_scalar_placeholders(self, doc: Document, values: Dict[str, str]) -> None:
        for paragraph in self._iter_all_paragraphs(doc):
            for run in paragraph.runs:
                for key, value in values.items():
                    token = "{{" + key + "}}"
                    if token in run.text:
                        run.text = run.text.replace(token, value)

    def _clone_repeating_group(self, doc: Document, rg, components: List[Dict[str, str]]) -> None:
        label_token = "{{" + rg.row_placeholders[0] + "}}"
        terms_token = "{{" + rg.row_placeholders[1] + "}}"

        table_row = self._find_template_table_row(doc, label_token)
        if table_row is not None:
            self._clone_table_row(table_row, label_token, terms_token, components)
            return

        para_pair = self._find_template_paragraph_pair(doc, label_token, terms_token)
        if para_pair is not None:
            self._clone_paragraph_pair(para_pair, label_token, terms_token, components)
            return

        raise ValueError(
            f"template for a repeating_group variant does not contain a template "
            f"row/block with placeholders {label_token}/{terms_token}"
        )

    @staticmethod
    def _find_template_table_row(doc: Document, label_token: str):
        for table in doc.tables:
            for row in table.rows:
                row_text = "".join(cell.text for cell in row.cells)
                if label_token in row_text:
                    return row
        return None

    @staticmethod
    def _clone_table_row(template_row, label_token: str, terms_token: str, components):
        template_tr = template_row._tr
        for entry in components:
            new_tr = copy.deepcopy(template_tr)
            DocTemplateEngine._replace_in_xml_element(new_tr, label_token, entry["label"])
            DocTemplateEngine._replace_in_xml_element(new_tr, terms_token, entry["terms"])
            template_tr.addprevious(new_tr)
        template_tr.getparent().remove(template_tr)

    @staticmethod
    def _find_template_paragraph_pair(doc: Document, label_token: str, terms_token: str):
        paragraphs = doc.paragraphs
        for i, p in enumerate(paragraphs):
            if label_token in p.text and i + 1 < len(paragraphs) and terms_token in paragraphs[i + 1].text:
                return paragraphs[i], paragraphs[i + 1]
        return None

    @staticmethod
    def _clone_paragraph_pair(pair, label_token: str, terms_token: str, components):
        label_p, terms_p = pair
        label_el, terms_el = label_p._p, terms_p._p
        for entry in components:
            new_label_el = copy.deepcopy(label_el)
            new_terms_el = copy.deepcopy(terms_el)
            DocTemplateEngine._replace_in_xml_element(new_label_el, label_token, entry["label"])
            DocTemplateEngine._replace_in_xml_element(new_terms_el, terms_token, entry["terms"])
            label_el.addprevious(new_label_el)
            label_el.addprevious(new_terms_el)
        label_el.getparent().remove(label_el)
        terms_el.getparent().remove(terms_el)

    @staticmethod
    def _replace_in_xml_element(element, token: str, value: str) -> None:
        # Text runs live at w:r/w:t under the (possibly deepcopy'd) w:tr/w:p element.
        from docx.oxml.ns import qn

        for t in element.iter(qn("w:t")):
            if t.text and token in t.text:
                t.text = t.text.replace(token, value)

    @staticmethod
    def _extract_full_text(doc: Document) -> str:
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    @staticmethod
    def count_pages(docx_path: Path) -> Optional[int]:
        """Real page count via a headless LibreOffice conversion to PDF, then
        counting pages with PyMuPDF/fitz - the same library `PDFExtractor`
        already depends on for the reverse direction. python-docx has no
        layout engine and can never itself know how a .docx paginates in
        real Word, so this is the only honest way to answer "is this one
        page?" (2026-09-14, explicit human instruction: fee agreements must
        never spill to a second page). Best-effort - returns None (never
        raises) if LibreOffice isn't available in this environment; callers
        must treat None as "unknown", never as "1 page, fine"."""
        import subprocess
        import tempfile

        try:
            import fitz  # PyMuPDF
        except ImportError:
            logger.warning("[083] PyMuPDF not available - cannot count pages")
            return None

        try:
            with tempfile.TemporaryDirectory() as tmp_out:
                result = subprocess.run(
                    [
                        "soffice", "--headless", "--convert-to", "pdf",
                        "--outdir", tmp_out, str(docx_path),
                    ],
                    capture_output=True, timeout=60, check=False,
                )
                if result.returncode != 0:
                    logger.warning(
                        f"[083] soffice conversion failed (rc={result.returncode}): "
                        f"{result.stderr!r}"
                    )
                    return None
                pdf_path = Path(tmp_out) / (docx_path.stem + ".pdf")
                if not pdf_path.exists():
                    logger.warning(f"[083] soffice produced no PDF for {docx_path}")
                    return None
                with fitz.open(str(pdf_path)) as pdf:
                    return pdf.page_count
        except (OSError, subprocess.SubprocessError) as e:
            logger.warning(f"[083] page count check failed: {e}")
            return None
